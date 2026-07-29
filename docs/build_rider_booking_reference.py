from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('RIDER_BOOKING_API_AND_DATABASE_REFERENCE.docx')

def font(run, size=11, bold=False, color=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def add_text(p, text, **kw):
    r = p.add_run(text); font(r, **kw); return r

def para(doc, text='', size=11, bold=False, color=None, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.25
    add_text(p, text, size=size, bold=bold, color=color); return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18 if level == 1 else 12); p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    add_text(p, text, size=16 if level == 1 else 13, bold=True, color='2E74B5' if level < 3 else '1F4D78')
    return p

def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_together = True
        add_text(p, value)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.style = 'Table Grid'; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''; shade(c, 'E8EEF5'); set_cell_margins(c)
        add_text(c.paragraphs[0], h, size=9, bold=True, color='1F4D78')
        if widths: c.width = Inches(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''; set_cell_margins(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_text(cells[i].paragraphs[0], str(value), size=8.5)
            if widths: cells[i].width = Inches(widths[i])
    return t

def page_break(doc): doc.add_page_break()

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)
styles = doc.styles
styles['Normal'].font.name = 'Calibri'; styles['Normal'].font.size = Pt(11)

# Header and footer
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(hp, 'Rider Booking Application | Technical Reference', size=8.5, color='666666')
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_text(fp, 'Internal engineering reference - generated 19 July 2026', size=8.5, color='666666')

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(72); p.paragraph_format.space_after = Pt(6)
add_text(p, 'RIDER BOOKING APPLICATION', size=24, bold=True, color='0B2545')
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(24); add_text(p, 'API, Flow, Database, and Operational Test Reference', size=15, color='2E74B5')
para(doc, 'Scope: ride booking, rider matching, live location, ride-start OTP, ride completion, invoices, authentication, and fare calculation.', size=11, after=18)
table(doc, ['Decision', 'Implemented behavior'], [
    ('Ride-start OTP expiry', 'Removed. A ride OTP has no clock-based expiry; it is accepted only for the assigned rider while the ride status is ACCEPTED.'),
    ('Stale location', 'Ignored when updated_at is 10 minutes old or older.'),
    ('Nearby matching', 'Returns all AVAILABLE riders with a fresh location inside a 10 km radius, sorted nearest first.'),
    ('Invoice creation', 'Automatically inserted when a STARTED ride is ended successfully. One invoice per ride, status PENDING.'),
], [1.65, 4.85])
heading(doc, 'How to use this guide', 1)
bullets(doc, ['Use the endpoint checklist to perform manual API verification.', 'Run the database script for a clean PostgreSQL setup, or apply the versioned Flyway migrations to an existing environment.', 'Protect secrets: move database, email, and SMS provider credentials out of application.properties before deployment.'])

page_break(doc)
heading(doc, '1. Ride lifecycle and business rules', 1)
table(doc, ['State / event', 'Rule and resulting action'], [
 ('REQUESTED', 'Booking stores the route, fare, and a secure 4-digit ride OTP; matching sends ride requests to eligible riders.'),
 ('ACCEPTED', 'First rider to win the atomic accept update is assigned. Other rider requests are cancelled.'),
 ('STARTED', 'Assigned rider supplies the correct OTP. There is no expiry timestamp; status must still be ACCEPTED.'),
 ('COMPLETED', 'Only a STARTED ride can end. The rider becomes AVAILABLE and a PENDING invoice is created in the same transaction.'),
 ('RIDER_CANCELLED', 'Assigned rider cancellation clears assignment, notifies the user, and attempts a rematch excluding already-notified riders.'),
 ('ADMIN_CANCELLED', 'User/admin cancellation cancels all rider request records and notifies the assigned rider if present.'),
], [1.55, 4.95])
heading(doc, 'Ride-start OTP: no expiry', 2)
para(doc, 'The removed expiry applies only to the booking ride-start OTP (RIDE_DETAILS.RIDE_OTP). The separate mobile-login OTP endpoints remain unchanged. A ride-start OTP is still rejected when the rider is not assigned, the ride is not ACCEPTED, or the OTP value is wrong. Cancellation or starting the ride makes the old state ineligible for another start.')
heading(doc, 'Stale-location rule', 2)
para(doc, 'Yes. The matching query uses updated_at > NOW() - INTERVAL \'10 minutes\'. Therefore a location exactly 10 minutes old, or older, is ignored. Freshness is assessed using database time, so application servers and the PostgreSQL server should use synchronized clocks.')
heading(doc, '10 km matching rule', 2)
para(doc, 'The Haversine formula calculates straight-line distance in kilometers. Matching requires all three conditions: the rider has a location newer than 10 minutes, rider status is AVAILABLE, and calculated distance is less than or equal to 10 km. There is no LIMIT, so all qualifying riders are returned and sorted by distance.')

page_break(doc)
heading(doc, '2. Endpoint checklist', 1)
para(doc, 'Base URL: http://localhost:8088. Unless listed as public, send Authorization: Bearer <JWT>. Current security configuration permits /register/**, /otp/**, /demo, /error, and /sendJWT without a JWT; all other HTTP endpoints require it.', size=10)
table(doc, ['Method & path', 'Input', 'Expected behavior'], [
 ('POST /register', 'User JSON', 'Creates user and starts email confirmation flow.'),
 ('GET /register', 'None', 'Returns registration page HTML.'),
 ('GET /register/confirmToken?token=', 'token', 'Confirms email token and returns status map.'),
 ('GET /register/verifyEmail?emailAddress=', 'emailAddress', 'Checks whether email exists/validates as implemented.'),
 ('POST /register/updateUser', 'Form User fields', 'Updates/registers user via legacy form flow.'),
 ('POST /login', '{username,password}', 'Authenticates and returns JWT.'),
 ('GET /sendJWT?username=', 'username', 'Public token generator; secure/restrict before production.'),
 ('POST /otp/send/sms', 'phone, userName parameters', 'Generates and sends mobile OTP.'),
 ('POST /otp/verify', 'key, otp, username parameters', 'Verifies mobile OTP and returns JWT when valid.'),
 ('GET /api/auth/validate', 'JWT', 'Returns authenticated username and valid=true.'),
], [1.75, 2.05, 2.7])

page_break(doc)
heading(doc, 'Ride endpoints', 1)
table(doc, ['Method & path', 'Input', 'Expected behavior'], [
 ('POST /ride/request', 'RideDetails JSON', 'Creates REQUESTED ride, permanent-in-state OTP, all eligible 10 km rider requests, and rider WebSocket notifications.'),
 ('POST /ride/accept?rideId=&riderId=', 'Query parameters', 'Atomically assigns the first accepting rider; cancels competing requests.'),
 ('POST /ride/verify-otp?rideId=&riderId=&otp=', 'Query parameters', 'Starts only an ACCEPTED ride for its assigned rider with correct OTP; no time-expiry check.'),
 ('POST /ride/end?rideId=&riderId=', 'Query parameters', 'Completes a STARTED ride, makes rider AVAILABLE, and creates a PENDING invoice.'),
 ('POST /ride/rider-cancel?rideId=&riderId=', 'Query parameters', 'Cancels assignment and attempts rematch.'),
 ('POST /ride/cancel?rideId=', 'Query parameter', 'Admin/user cancellation; cancels all request rows.'),
 ('POST /rider/fetch/livelocation', 'RiderLocation JSON', 'Publishes location to /topic/ride/{rideId}.'),
 ('POST /rider/save/location', 'RiderLocation JSON', 'Upserts rider location and refreshes updated_at.'),
 ('POST /rider', 'RiderDetails JSON', 'Creates a rider in RIDER_DETAILS. Status defaults to AVAILABLE.'),
 ('GET /rider/nearby?lat=&lon=', 'Query parameters', 'Returns rider IDs for all fresh, AVAILABLE riders within 10 km.'),
 ('GET /api/fare/calculate?sLat=&sLon=&dLat=&dLon=', 'Query parameters', 'Returns distance/fare display string; fare calculation is presently a placeholder.'),
 ('GET /demo', 'None', 'Triggers demo rider details service and returns home page.'),
], [1.85, 1.95, 2.7])
heading(doc, 'Payloads to test', 2)
para(doc, 'POST /ride/request', bold=True)
para(doc, '{"pickUpLatitude":19.0760,"pickUpLongitude":72.8777,"dropLatitude":19.1100,"dropLongitude":72.9000,"fare":125.00,"adminId":1}', size=9, color='1F4D78')
para(doc, 'POST /rider/save/location', bold=True)
para(doc, '{"rideId":0,"riderId":12,"latitude":19.0760,"longitude":72.8777}', size=9, color='1F4D78')

page_break(doc)
heading(doc, '3. WebSocket functionality checklist', 1)
table(doc, ['Connection / destination', 'Purpose', 'Test'], [
 ('CONNECT /ws (SockJS + STOMP)', 'JWT interceptor authenticates client CONNECT frames.', 'Connect with valid JWT; confirm invalid/missing token fails according to interceptor.'),
 ('Subscribe /topic/rider/{riderId}', 'Receives RIDE_REQUEST, RIDE_CANCELLED, and RIDE_CANCELLED_BY_USER events.', 'Create then accept/cancel a booking and confirm event payload.'),
 ('Subscribe /topic/user/{userId}', 'Receives accepted, started, cancelled, and completed events.', 'Follow a complete booking through acceptance, start, and end.'),
 ('Subscribe /topic/ride/{rideId}', 'Receives location, RIDE_STARTED, and RIDE_COMPLETED events.', 'Post live location and progress a ride.'),
], [2.1, 2.2, 2.2])
heading(doc, 'Functional test sequence', 2)
bullets(doc, ['Save two AVAILABLE riders within 10 km and one beyond 10 km; call GET /rider/nearby and verify only the two close riders return.', 'Set one close rider updated_at to more than 10 minutes ago; verify that rider no longer returns.', 'Request a ride, accept it, wait more than ten minutes if desired, then verify the correct ride OTP still starts the ride.', 'End the STARTED ride and query INVOICES; confirm exactly one row with the ride ID, amount equal to fare, and status PENDING.', 'Call /ride/end a second time; it must fail because the ride is no longer STARTED.'])

page_break(doc)
heading(doc, '4. Automatic invoice behavior', 1)
para(doc, 'An invoice is created automatically by RideService.endRide after it validates that (1) the caller rider is assigned and (2) the ride is STARTED. In the transaction, it changes RIDE_DETAILS.RIDE_STATUS to COMPLETED, changes the rider to AVAILABLE, inserts INVOICES with amount = ride fare and status = PENDING, then publishes completion events. It does not charge a payment method, calculate tax, apply discounts, or settle funds; those are future payment-processing responsibilities.')
table(doc, ['Invoice field', 'Source / meaning'], [
 ('INVOICE_ID', 'Generated database primary key.'), ('RIDE_ID', 'Completed ride ID; UNIQUE constraint enforces one invoice per ride.'), ('RIDER_ID', 'Assigned rider who completed the ride.'), ('USER_ID', 'Ride ADMIN_ID (the booking user in current code).'), ('AMOUNT', 'Current ride fare; no tax/discount/commission logic yet.'), ('CREATED_AT', 'Time of successful end operation.'), ('STATUS', 'PENDING; intended for downstream payment/settlement.'),
], [1.5, 5.0])
heading(doc, 'Recommended controls before production', 2)
bullets(doc, ['Use an authenticated identity rather than caller-supplied adminId/riderId parameters to enforce ownership.', 'Add a payment workflow that changes PENDING to SETTLED or REFUNDED and records gateway references.', 'Rate-limit OTP attempts and redact OTP values from notifications/logs where not required.', 'Move SMTP, SMS, and database credentials into secret-managed environment variables.'])

page_break(doc)
heading(doc, '5. Database reference', 1)
para(doc, 'The companion script DATABASE_SETUP_AND_REFERENCE.sql is the detailed PostgreSQL setup script. It creates the tables, constraints, indexes, and the exact 10 km fresh-location reference query.')
table(doc, ['Table', 'Purpose', 'Key columns'], [
 ('USERS', 'Application users and account verification data.', 'ID, USER_NAME, EMAIL, ENABLED, TOKEN, mobile fields'),
 ('RIDER_DETAILS', 'Rider profile and availability.', 'ID, STATUS: AVAILABLE / ON_TRIP / OFFLINE'),
 ('RIDER_LOCATION', 'One current position per rider.', 'RIDER_ID PK, coordinates, UPDATED_AT'),
 ('RIDE_DETAILS', 'Ride lifecycle record.', 'RIDE_ID, route, fare, RIDER_ID, ADMIN_ID, RIDE_OTP, RIDE_STATUS'),
 ('RIDE_REQUEST', 'Per-rider matching/notification record.', 'RIDE_ID, RIDER_ID, STATUS, CREATED_AT'),
 ('INVOICES', 'One completion invoice per ride.', 'RIDE_ID UNIQUE, amount, created_at, status'),
], [1.4, 2.45, 2.65])
heading(doc, 'Migration order', 2)
bullets(doc, ['V1 creates RIDE_REQUEST.', 'V2 adds rider status (legacy migration names RIDER_DEMO; the current JPA entity maps RIDER_DETAILS, so use the companion setup script for a clean schema and verify the deployed table name).', 'V3 creates INVOICES.', 'V4 added OTP_EXPIRES_AT.', 'V5 removes OTP_EXPIRES_AT and its index because ride-start OTPs no longer expire by time.'])
heading(doc, 'Database validation queries', 2)
para(doc, 'Fresh eligible rider count: SELECT COUNT(*) FROM RIDER_LOCATION rl JOIN RIDER_DETAILS rd ON rd.ID=rl.RIDER_ID WHERE rl.UPDATED_AT > CURRENT_TIMESTAMP - INTERVAL \'10 minutes\' AND rd.STATUS=\'AVAILABLE\';', size=9, color='1F4D78')
para(doc, 'Invoice check: SELECT * FROM INVOICES WHERE RIDE_ID = :rideId;', size=9, color='1F4D78')

page_break(doc)
heading(doc, '6. Deployment and review checklist', 1)
bullets(doc, ['Back up the database before applying the V5 migration to any shared environment.', 'Apply V5__remove_ride_otp_expiry.sql after existing V1-V4 migrations. Confirm OTP_EXPIRES_AT is absent from RIDE_DETAILS.', 'Run mvn test. The project test context can start without a reachable local database, but endpoint and SQL behavior must be validated against PostgreSQL separately.', 'Run the manual sequence in section 3 with a real JWT and WebSocket client.', 'Confirm the deployed table name is consistently RIDER_DETAILS across JPA mappings, migrations, and the actual database.', 'Confirm system clocks are synchronized; stale-location evaluation depends on PostgreSQL CURRENT_TIMESTAMP.', 'Do not expose /sendJWT publicly in production.'])
heading(doc, 'Known implementation notes', 2)
bullets(doc, ['The nearby endpoint returns rider IDs only. Add a response DTO with name, coordinates, and calculated distance if the client needs display data.', 'The fare endpoint currently formats a placeholder rather than applying a final pricing formula.', 'Email/SMS secrets appear in the local configuration file and must be rotated if they are real credentials.', 'Current endpoint authorization is broad; use role/ownership authorization for rider and user actions.'])

doc.save(OUT)
print(OUT)
